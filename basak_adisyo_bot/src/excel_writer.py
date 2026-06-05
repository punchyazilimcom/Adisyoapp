"""Mevcut .xlsm şablonunu açıp SADECE izinli hücrelere yazar.

Kurallar:
- keep_vba=True ZORUNLU (makrolar korunur).
- Hedef sayfa: işlenen ayın/günün adı, ör. "HAZİRAN (6)".
- ALLOWED_CELLS dışına TEK bir hücreye bile yazılmaz (guard).
- Sadece ``.value`` atanır → hücre formatı/biçimi bozulmaz.
- Sayfa yoksa logla, o şubeyi atla, ÇÖKME.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

log = logging.getLogger("basak.excel")

# Bot'un yazmasına İZİN VERİLEN tek hücreler. Bunun dışına asla yazılmaz.
ALLOWED_CELLS: frozenset[str] = frozenset(
    {
        # CİRO (adet, L sütunu)
        "L7",   # KÜÇÜK AYRAN
        "L11",  # ÇAY
        "L17",  # KIR PİDELERİ (Kuşbaşılı hariç)
        "L19",  # KUTU İÇECEKLER
        # MENÜLER (adet, E sütunu)
        "E16",  # KURYE MENÜ
        "E18",  # KUTU İÇECEK MENÜ
        "E20",  # BÜYÜK AYRAN MENÜ
        # ÖDEME (tutar)
        "R9", "Q10",   # YS Online net / indirim
        "R11", "Q12",  # Trendyol Online net / indirim
        "R13", "Q14",  # Getir Online net / indirim
        "R15", "Q16",  # Migros Online net / indirim
    }
)

# Türkçe ay adları (büyük harf) — şablon sayfa adı bu önekle başlar.
TURKISH_MONTHS: dict[int, str] = {
    1: "OCAK",
    2: "ŞUBAT",
    3: "MART",
    4: "NİSAN",
    5: "MAYIS",
    6: "HAZİRAN",
    7: "TEMMUZ",
    8: "AĞUSTOS",
    9: "EYLÜL",
    10: "EKİM",
    11: "KASIM",
    12: "ARALIK",
}


def sheet_name_for(month: int, day: int) -> str:
    """Ör: (6, 6) → 'HAZİRAN (6)'."""
    return f"{TURKISH_MONTHS[month]} ({day})"


def write_day_cells(
    excel_path: str | Path,
    month: int,
    day: int,
    cell_values: dict[str, Any],
    *,
    dry_run: bool = False,
) -> bool:
    """Şablonu aç, hedef sayfaya izinli hücreleri yaz, kaydet.

    Dönüş: yazma başarılıysa True; sayfa/dosya yoksa False (çökmeden).
    """
    excel_path = Path(excel_path)
    target_sheet = sheet_name_for(month, day)

    # Guard: izinsiz hücre var mı? (programlama hatasına karşı erken patla)
    illegal = set(cell_values) - ALLOWED_CELLS
    if illegal:
        raise ValueError(f"İzinsiz hücrelere yazma girişimi: {sorted(illegal)}")

    if dry_run:
        log.info("[dry-run] '%s' sayfasına yazılmayacak (sadece hesaplandı).", target_sheet)
        return True

    if not excel_path.exists():
        log.error("Excel bulunamadı, şube atlanıyor: %s", excel_path)
        return False

    wb = load_workbook(excel_path, keep_vba=True)
    try:
        if target_sheet not in wb.sheetnames:
            log.error(
                "'%s' sayfası şablonda yok (%s). Şube atlanıyor.",
                target_sheet,
                excel_path.name,
            )
            return False

        ws = wb[target_sheet]
        for cell, value in cell_values.items():
            if cell not in ALLOWED_CELLS:  # ikinci güvenlik katmanı
                raise ValueError(f"İzinsiz hücre: {cell}")
            if value is None:
                continue
            ws[cell].value = value  # SADECE değer; biçim/format korunur

        wb.save(excel_path)
        log.info("Yazıldı → %s :: %s", excel_path.name, target_sheet)
        return True
    finally:
        wb.close()
